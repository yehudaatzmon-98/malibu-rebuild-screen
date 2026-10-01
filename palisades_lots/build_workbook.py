import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule
import os
HERE = os.path.dirname(os.path.abspath(__file__)); DATA = os.path.join(HERE, "data")
cur='$#,##0;($#,##0);-'; pct='0.0%;(0.0%);-'; num='#,##0;(#,##0);-'
# Neighborhood sale $/sf: maintained by update_comps.py (applied_psf; override wins)
ex = []
for b in csv.DictReader(open(os.path.join(DATA, "nbhd_base.csv"))):
    val = float(b["override_psf"]) if b["override_psf"] else float(b["applied_psf"])
    how = b["status"] + (f"; {b['rule_n']} qualifying comps, median ${int(float(b['rule_psf'])):,}" if b["rule_psf"] else "; no qualifying comps yet")
    ex.append((b["code"], val, f"{b['note']}  [{how}; updated {b['updated']}]"))
COMPS_LOG = [r for r in csv.DictReader(open(os.path.join(DATA, "comps_log.csv"))) if r["qualifies"] == "Y"]
ASKING = list(csv.DictReader(open(os.path.join(DATA, "asking_newbuilds.csv"))))

F = "Arial"
NAVY = "1F3864"
hf = Font(name=F, bold=True, color="FFFFFF", size=10)
body = Font(name=F, size=10)
small = Font(name=F, size=9, color="595959")
blue = Font(name=F, size=10, color="0000FF")
bold = Font(name=F, size=10, bold=True)
title = Font(name=F, size=15, bold=True, color=NAVY)
sub = Font(name=F, size=11, bold=True, color=NAVY)
yel = PatternFill("solid", start_color="FFF2CC")
thin = Side(style="thin", color="D9D9D9")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
GROUP = {"lot": "1F3864", "house": "2E75B6", "cost": "7F6000", "ret": "375623", "scen": "7030A0", "calc": "808080"}
tint = {"lot": "DDE4F0", "house": "DEEBF7", "cost": "FFF2CC", "ret": "E2EFDA", "scen": "E4DFEC", "calc": "EDEDED"}
money0 = '$#,##0;($#,##0);"-"'
psf = '$#,##0"/sf"'
NB = {"RIV": "Riviera", "HUNT": "Huntington", "RUSTIC": "Rustic Canyon", "BLUFF": "Bluffs / El Medio", "VDLP": "Via de la Paz",
      "ALPHA": "Alphabet / Village", "MK": "Marquez Knolls", "UPS": "Upper Las Lomas / Las Pulgas", "RIDGE": "Ridgeview",
      "CAST": "Castellammare", "ENC": "The Enclave", "HIGH": "Highlands", "SUMMIT": "The Summit"}

wb = Workbook()

# ================= Assumptions =================
a = wb.active; a.title = "Assumptions"
a['A1'] = "Assumptions — change the yellow cells and the whole workbook updates"; a['A1'].font = title
a['A3'], a['B3'], a['C3'] = "Input", "Value", "Where it comes from"
for c in a[3]: c.font = hf; c.fill = PatternFill("solid", start_color=NAVY)
inp=[
("Construction cost ($/sf, fixed)",700,cur,"Tal, 9/29 call: run one model at $700/sf fixed. Contractor Tal's quote; Tal says local builders are at ~$600 (secondhand). Any profit-share deal with the contractor would be upside."),
("Design, permits & fees — lots WITHOUT approved plans ($/sf)",75,cur,"Placeholder: architect, engineering, plan check, fees, insurance, PM."),
("Reserve — lots WITH city-approved plans (% of asking price)",0.05,pct,"Tal, 9/29 call: 5% of list price for minor approvals and GC items when plans are already approved."),
("Construction loan (% of cost)",0.65,pct,"Placeholder. Confirm with lender."),
("Loan rate",0.10,pct,"Placeholder."),
("Avg % of loan drawn",0.60,pct,"Placeholder for the draw schedule."),
("Loan points",0.015,pct,"Placeholder."),
("Property tax (% of land / yr)",0.012,pct,"~1.2% of purchase price (reassessed on transfer)."),
("Sale commission",0.05,pct,"Placeholder."),
("City + county transfer tax",0.0056,pct,"LA City 0.45% + County 0.11%."),
("ULA tier 1 threshold",5400000,cur,"Measure ULA, current threshold (LA City Clerk, Measure TE text)."),
("ULA tier 1 rate",0.04,pct,"Measure ULA."),
("ULA tier 2 threshold",10900000,cur,"Measure ULA, current threshold."),
("ULA tier 2 rate",0.055,pct,"Measure ULA. The Nov 2026 ballot exemption may not cover builder resales, so assumed to apply."),
("Fast-track size (x old house)",1.10,'0.00x',"LA EO1 like-for-like rebuild cap (110%)."),
("Build-to-zoning size (x lot area)",0.40,'0.00x',"Rough proxy only; an architect must confirm per lot."),
("Build-to-zoning max sf",8000,num,"Cap to keep spec homes a sellable size."),
("Fast-track / approved-plans project length (yrs)",2,num,"RTI lots start right away."),
("Build-to-zoning project length (yrs)",3,num,"+1 year for design and plan check."),
("Target margin (profit ÷ total cost)",0.15,pct,"Typical spec-build hurdle."),
("WORST case: premium over today's comps",0.0,pct,"Today's closed sales, no new-build premium. Fire history (Malibu, Altadena) shows new builds selling flat or below in years 1–3."),
("MEDIUM case: premium over today's comps",0.10,pct,"Tal, 9/29: new construction should beat today's comps by 10–15%."),
("BEST case: premium over today's comps",0.20,pct,"Recovery case: neighborhood rebuilt and demand returns."),
("Scenario used in the cost breakdown (Worst / Medium / Best)","Medium",'@',"Type Worst, Medium or Best."),
]
N = {}
for i, (k, v, fmt, src) in enumerate(inp, 4):
    a.cell(row=i, column=1, value=k).font = body
    c = a.cell(row=i, column=2, value=v); c.number_format = fmt
    if isinstance(v, str) and v.startswith('='): c.font = body
    else: c.font = blue; c.fill = yel
    a.cell(row=i, column=3, value=src).font = small
    N[k] = f"Assumptions!$B${i}"
r0 = len(inp) + 6
a.cell(row=r0, column=1, value="Sale price per sq ft by neighborhood (new 4,500–7,000 sf home)").font = sub
for j, h in enumerate(["Code", "Sale $/sf", "Neighborhood · strength of evidence"], 1):
    c = a.cell(row=r0 + 1, column=j, value=h); c.font = hf; c.fill = PatternFill("solid", start_color=NAVY)
for i, (k, v, n) in enumerate(ex, r0 + 2):
    a.cell(row=i, column=1, value=k).font = body
    c = a.cell(row=i, column=2, value=v); c.font = blue; c.fill = yel; c.number_format = cur
    a.cell(row=i, column=3, value=n).font = small
EXR = f"Assumptions!$A${r0+2}:$A${r0+1+len(ex)}"; EXV = f"Assumptions!$B${r0+2}:$B${r0+1+len(ex)}"
a.column_dimensions['A'].width = 52; a.column_dimensions['B'].width = 14; a.column_dimensions['C'].width = 100
a.freeze_panes = "A4"

# ================= All Lots =================
lt = wb.create_sheet("All Lots", 0)
cols = [  # (header, group, width, fmt)
 ("Rank", "lot", 6, None), ("Address", "lot", 26, None), ("Neighborhood", "lot", 18, None), ("Asking price", "lot", 13, money0),
 ("Lot size (sf)", "lot", 10, num), ("Old house (sf)", "lot", 10, num), ("City-approved / claimed plans (sf)", "lot", 13, num), ("City-approved plans?", "lot", 9, None), ("Notes", "lot", 40, None),
 ("New house size (sf)", "house", 11, num), ("How sized", "house", 17, None), ("Project length (yrs)", "house", 9, num),
 ("Sale price per sf", "house", 11, psf), ("Sale price", "house", 13, money0),
 ("Land", "cost", 13, money0), ("Construction", "cost", 13, money0), ("Design, permits & fees (or 5% reserve)", "cost", 13, money0),
 ("Loan interest & points", "cost", 12, money0), ("Property tax", "cost", 11, money0), ("Commission & transfer tax", "cost", 12, money0),
 ("Mansion tax (ULA)", "cost", 12, money0), ("TOTAL COST", "cost", 14, money0),
 ("Profit", "ret", 13, money0), ("Margin (profit ÷ cost)", "ret", 9, '0%;(0%);"-"'), ("Most we should pay for the lot", "ret", 14, money0), ("Asking vs. most we should pay", "ret", 11, '0%'),
 ("Worst: sale $/sf", "scen", 10, psf), ("Worst: profit", "scen", 12, money0), ("Medium: sale $/sf", "scen", 10, psf), ("Medium: profit", "scen", 12, money0), ("Best: sale $/sf", "scen", 10, psf), ("Best: profit", "scen", 12, money0),
 ("Nbhd code", "calc", 8, None), ("Fast-path size", "calc", 9, num), ("Build-to-zoning size", "calc", 9, num),
 ("Profit: fast path", "calc", 12, money0), ("Profit: build to zoning", "calc", 12, money0), ("Loan amount", "calc", 12, money0),
]
col = {h: L(i) for i, (h, _, _, _) in enumerate(cols, 1)}
# group banner row 1, headers row 2
banners = [("THE LOT", "lot"), ("THE HOUSE WE BUILD", "house"), ("COSTS", "cost"), ("RETURNS (scenario chosen on Assumptions)", "ret"), ("PROFIT IN EACH SCENARIO", "scen"), ("BEHIND THE SCENES (math helpers)", "calc")]
for name, g in banners:
    idx = [i for i, (_, gg, _, _) in enumerate(cols, 1) if gg == g]
    lt.merge_cells(start_row=1, start_column=idx[0], end_row=1, end_column=idx[-1])
    c = lt.cell(row=1, column=idx[0], value=name); c.font = hf; c.fill = PatternFill("solid", start_color=GROUP[g]); c.alignment = center
for i, (h, g, w, _) in enumerate(cols, 1):
    c = lt.cell(row=2, column=i, value=h); c.font = Font(name=F, bold=True, size=9, color="000000")
    c.fill = PatternFill("solid", start_color=tint[g]); c.alignment = center
    lt.column_dimensions[L(i)].width = w
lt.row_dimensions[2].height = 42

order = open(os.path.join(DATA, 'order.txt')).read().split("\n")
raw = {r[0]: r for r in csv.reader(open(os.path.join(DATA, 'lots.csv'))) if r and not r[0].startswith('#')}
n = N
bc = n['Construction cost ($/sf, fixed)']; sc = n['Design, permits & fees — lots WITHOUT approved plans ($/sf)']; RES = n['Reserve — lots WITH city-approved plans (% of asking price)']
LTC, RATE, DRAW, PTS, PTAX = n['Construction loan (% of cost)'], n['Loan rate'], n['Avg % of loan drawn'], n['Loan points'], n['Property tax (% of land / yr)']
COMM, XF = n['Sale commission'], n['City + county transfer tax']
T1, R1, T2, R2 = n['ULA tier 1 threshold'], n['ULA tier 1 rate'], n['ULA tier 2 threshold'], n['ULA tier 2 rate']
M = n['Target margin (profit ÷ total cost)']
PW, PM, PB, SCN = n["WORST case: premium over today's comps"], n["MEDIUM case: premium over today's comps"], n["BEST case: premium over today's comps"], n['Scenario used in the cost breakdown (Worst / Medium / Best)']
PSEL = f'IF(UPPER({SCN})="WORST",{PW},IF(UPPER({SCN})="BEST",{PB},{PM}))'

def ula(sale): return f"IF({sale}>={T2},{R2},IF({sale}>={T1},{R1},0))"
def soft(sf, ask, appr): return f"IF({appr}=\"Yes\",{RES}*({ask}),({sf})*{sc})"
def profit_expr(sf, yrs, ask, exps, appr):
    sale = f"({sf})*{exps}"
    base = f"(({ask})+({sf})*{bc}+{soft(sf, ask, appr)})"
    cost = f"({base}*(1+{LTC}*({RATE}*({yrs})*{DRAW}+{PTS}))+({ask})*{PTAX}*({yrs})+{sale}*({COMM}+{XF}+{ula(sale)}))"
    return f"{sale}-{cost}"

for r, addr in enumerate(order, 3):
    a_, code, ask, prior, plan, lotsf, flag, appr = raw[addr]
    c = lambda h: f"{col[h]}{r}"
    vals = {
     "Rank": r - 2, "Address": addr, "Neighborhood": NB[code], "Asking price": float(ask), "Lot size (sf)": float(lotsf),
     "Old house (sf)": float(prior), "City-approved / claimed plans (sf)": float(plan) if plan else None, "City-approved plans?": "Yes" if appr == "Y" else "No",
     "Notes": flag, "Nbhd code": code,
    }
    for h, v in vals.items(): lt[c(h)] = v
    AP = c("City-approved plans?")
    ya = n['Fast-track / approved-plans project length (yrs)']
    yb = f"IF({c('Build-to-zoning size')}>{c('Fast-path size')},{n['Build-to-zoning project length (yrs)']},{ya})"
    comp = f"INDEX({EXV},MATCH({c('Nbhd code')},{EXR},0))"
    lt[c("Fast-path size")] = f"=IF({c('City-approved / claimed plans (sf)')}>0,{c('City-approved / claimed plans (sf)')},ROUND({c('Old house (sf)')}*{n['Fast-track size (x old house)']},0))"
    lt[c("Build-to-zoning size")] = f"=IF({AP}=\"Yes\",{c('Fast-path size')},MAX({c('Fast-path size')},MIN({n['Build-to-zoning max sf']},ROUND({c('Lot size (sf)')}*{n['Build-to-zoning size (x lot area)']},0))))"
    lt[c("Worst: sale $/sf")] = f"={comp}*(1+{PW})"
    lt[c("Medium: sale $/sf")] = f"={comp}*(1+{PM})"
    lt[c("Best: sale $/sf")] = f"={comp}*(1+{PB})"
    lt[c("Sale price per sf")] = f"={comp}*(1+{PSEL})"
    lt[c("Profit: fast path")] = "=" + profit_expr(c('Fast-path size'), ya, c('Asking price'), c('Medium: sale $/sf'), AP)
    lt[c("Profit: build to zoning")] = "=" + profit_expr(c('Build-to-zoning size'), yb, c('Asking price'), c('Medium: sale $/sf'), AP)
    big = f"{c('Profit: build to zoning')}>{c('Profit: fast path')}"
    lt[c("How sized")] = f'=IF({big},"Zoning estimate",IF({AP}="Yes","Approved plans",IF({c("City-approved / claimed plans (sf)")}>0,"Listing claim","110% old house")))'
    lt[c("New house size (sf)")] = f"=IF({big},{c('Build-to-zoning size')},{c('Fast-path size')})"
    lt[c("Project length (yrs)")] = f"=IF({big},{yb},{ya})"
    S_, Y_ = c('New house size (sf)'), c('Project length (yrs)')
    for sc_ in ("Worst", "Medium", "Best"):
        lt[c(f"{sc_}: profit")] = "=" + profit_expr(S_, Y_, c('Asking price'), c(f'{sc_}: sale $/sf'), AP)
    lt[c("Sale price")] = f"={S_}*{c('Sale price per sf')}"
    lt[c("Land")] = f"={c('Asking price')}"
    lt[c("Construction")] = f"={S_}*{bc}"
    lt[c("Design, permits & fees (or 5% reserve)")] = "=" + soft(S_, c('Asking price'), AP)
    lt[c("Loan amount")] = f"={LTC}*({c('Land')}+{c('Construction')}+{c('Design, permits & fees (or 5% reserve)')})"
    lt[c("Loan interest & points")] = f"={c('Loan amount')}*({RATE}*{Y_}*{DRAW}+{PTS})"
    lt[c("Property tax")] = f"={c('Land')}*{PTAX}*{Y_}"
    lt[c("Commission & transfer tax")] = f"={c('Sale price')}*({COMM}+{XF})"
    lt[c("Mansion tax (ULA)")] = f"={c('Sale price')}*{ula(c('Sale price'))}"
    lt[c("TOTAL COST")] = f"=SUM({c('Land')}:{c('Mansion tax (ULA)')})"
    lt[c("Profit")] = f"={c('Sale price')}-{c('TOTAL COST')}"
    lt[c("Margin (profit ÷ cost)")] = f"=IF({c('TOTAL COST')}=0,0,{c('Profit')}/{c('TOTAL COST')})"
    sale = c('Sale price'); y = Y_
    K = f"({S_}*({bc}+IF({AP}=\"Yes\",0,{sc})))"
    rr = f"IF({AP}=\"Yes\",{RES},0)"
    f_ = f"({LTC}*({RATE}*{y}*{DRAW}+{PTS}))"
    sell = f"{sale}*({COMM}+{XF}+{ula(sale)})"
    lt[c("Most we should pay for the lot")] = f"=MAX(0,({sale}/(1+{M})-{K}*(1+{f_})-{sell})/((1+{rr})*(1+{f_})+{PTAX}*{y}))"
    lt[c("Asking vs. most we should pay")] = f'=IF({c("Most we should pay for the lot")}<=0,"no price works",{c("Asking price")}/{c("Most we should pay for the lot")})'
    for i, (h, g, w, fmt) in enumerate(cols, 1):
        cell = lt.cell(row=r, column=i)
        cell.font = blue if h in ("Asking price", "Lot size (sf)", "Old house (sf)", "City-approved / claimed plans (sf)", "City-approved plans?") else (bold if h in ("TOTAL COST", "Profit") else body)
        if fmt: cell.number_format = fmt
        cell.border = Border(bottom=thin)
        if h == "Notes": cell.alignment = Alignment(wrap_text=True, vertical="top")
last = len(order) + 2
rng = f"A3:{col['Nbhd code']}{last}"
lt.conditional_formatting.add(f"{col['Asking vs. most we should pay']}3:{col['Asking vs. most we should pay']}{last}",
    FormulaRule(formula=[f'AND(ISNUMBER({col["Asking vs. most we should pay"]}3),{col["Asking vs. most we should pay"]}3<=1)'], fill=PatternFill("solid", start_color="C6EFCE")))
lt.conditional_formatting.add(f"{col['Asking vs. most we should pay']}3:{col['Asking vs. most we should pay']}{last}",
    FormulaRule(formula=[f'OR(NOT(ISNUMBER({col["Asking vs. most we should pay"]}3)),{col["Asking vs. most we should pay"]}3>1)'], fill=PatternFill("solid", start_color="F8CBAD")))
lt.freeze_panes = "C3"; lt.auto_filter.ref = f"A2:{col['Loan amount']}{last}"
for h in ("Nbhd code", "Fast-path size", "Build-to-zoning size", "Profit: fast path", "Profit: build to zoning", "Loan amount"):
    lt.column_dimensions[col[h]].outlineLevel = 1
lt.sheet_properties.outlinePr.summaryRight = False

# ================= Deal Breakdown (shortlist) =================
db = wb.create_sheet("Deal Breakdown", 0)
short = open(os.path.join(DATA, 'short.txt')).read().split("\n")
db['A1'] = "Palisades lots — deal breakdown for the shortlist"; db['A1'].font = title
db['A2'] = f'="Construction at $"&TEXT({bc},"#,##0")&"/sf fixed. Cost breakdown shows the "&{SCN}&" case; all three cases are at the bottom. Change inputs on the Assumptions tab."'
db['A2'].font = small
db['A3'] = "Every number is pulled live from the All Lots tab. Sale prices start from closed sales by neighborhood; Worst = today's comps, Medium = +10%, Best = +20%."; db['A3'].font = small
lines = [("THE LOT", None, "lot"), ("Neighborhood", "Neighborhood", None), ("Asking price", "Asking price", money0), ("Lot size (sf)", "Lot size (sf)", num),
 ("City-approved plans?", "City-approved plans?", None),
 ("THE HOUSE", None, "house"), ("New house size (sf)", "New house size (sf)", num), ("How sized", "How sized", None),
 ("Sale price per sf", "Sale price per sf", psf), ("Projected sale price", "Sale price", money0),
 ("COSTS", None, "cost"), ("Land", "Land", money0), ("Construction", "Construction", money0), ("Design, permits & fees (or 5% reserve)", "Design, permits & fees (or 5% reserve)", money0),
 ("Loan interest & points", "Loan interest & points", money0), ("Property tax", "Property tax", money0), ("Commission & transfer tax", "Commission & transfer tax", money0),
 ("Mansion tax (ULA)", "Mansion tax (ULA)", money0), ("Total cost", "TOTAL COST", money0),
 ("RETURNS", None, "ret"), ("Profit", "Profit", money0), ("Margin (profit ÷ cost)", "Margin (profit ÷ cost)", '0%;(0%);"-"'),
 ("Most we should pay for the lot", "Most we should pay for the lot", money0), ("Asking vs. most we should pay", "Asking vs. most we should pay", '0%'),
 ("PROFIT IN EACH SCENARIO", None, "scen"),
 ("Worst: sale price per sf (today's comps)", "Worst: sale $/sf", psf), ("Worst: profit", "Worst: profit", money0),
 ("Medium: sale price per sf (+10%)", "Medium: sale $/sf", psf), ("Medium: profit", "Medium: profit", money0),
 ("Best: sale price per sf (+20%)", "Best: sale $/sf", psf), ("Best: profit", "Best: profit", money0)]
hr = 5
db.cell(row=hr, column=1, value="").fill = PatternFill("solid", start_color=NAVY)
for j, addr in enumerate(short, 2):
    c = db.cell(row=hr, column=j, value=addr.replace(" St", "").replace(" Ave", "").replace(" Dr", "").replace(" Rd", "").replace(" Blvd", "").replace(" Pl", ""))
    c.font = hf; c.fill = PatternFill("solid", start_color=NAVY); c.alignment = center
    db.column_dimensions[L(j)].width = 16
db.column_dimensions['A'].width = 36; db.row_dimensions[hr].height = 32
BOLD = ("Total cost", "Profit", "Projected sale price", "Worst: profit", "Medium: profit", "Best: profit")
for i, (label, src, fmt) in enumerate(lines, hr + 1):
    if src is None:
        for j in range(1, len(short) + 2):
            cc = db.cell(row=i, column=j); cc.fill = PatternFill("solid", start_color=tint[fmt])
        db.cell(row=i, column=1, value=label).font = Font(name=F, bold=True, size=9, color=GROUP[fmt])
        continue
    fnt = bold if label in BOLD else body
    lab = db.cell(row=i, column=1, value=label); lab.font = fnt
    for j, addr in enumerate(short, 2):
        f = f"=INDEX('All Lots'!${col[src]}$3:${col[src]}${last},MATCH(\"{addr}\",'All Lots'!$B$3:$B${last},0))"
        cc = db.cell(row=i, column=j, value=f); cc.font = fnt
        if fmt: cc.number_format = fmt
        cc.alignment = Alignment(horizontal="right" if fmt else "center")
        cc.border = Border(bottom=thin)
        if label in ("Total cost", "Profit"):
            cc.border = Border(top=Side(style="thin", color="000000"), bottom=thin)
    if label == "Total cost": lab.border = Border(top=Side(style="thin", color="000000"))
ar = hr + len(lines)
rowAsk = hr + 1 + [l[0] for l in lines].index("Asking vs. most we should pay")
db.conditional_formatting.add(f"B{rowAsk}:{L(len(short)+1)}{rowAsk}", FormulaRule(formula=[f"AND(ISNUMBER(B{rowAsk}),B{rowAsk}<=1)"], fill=PatternFill("solid", start_color="C6EFCE")))
db.conditional_formatting.add(f"B{rowAsk}:{L(len(short)+1)}{rowAsk}", FormulaRule(formula=[f"OR(NOT(ISNUMBER(B{rowAsk})),B{rowAsk}>1)"], fill=PatternFill("solid", start_color="F8CBAD")))
for lbl in ("Worst: profit", "Medium: profit", "Best: profit", "Profit"):
    rw = hr + 1 + [l[0] for l in lines].index(lbl)
    db.conditional_formatting.add(f"B{rw}:{L(len(short)+1)}{rw}", FormulaRule(formula=[f"B{rw}<0"], font=Font(color="C00000")))
notes = [
 "How to read it",
 "• Total cost = everything from buying the lot to selling the house, including loan costs, the realtor commission and LA's mansion tax (ULA) at sale.",
 "• Profit = sale price − total cost. 'Most we should pay' = the highest lot price that still earns the target margin (15%). Green = the lot works at asking; red = we'd need to negotiate down.",
 "• Construction is $700/sf fixed. Lots with city-approved plans (checked in LADBS 9/28/2026) use a 5% of asking price reserve instead of design and permit costs.",
 "• House size: city-approved plans where they exist; otherwise 110% of the old house or a rough zoning estimate an architect must confirm.",
 "• Placeholders to confirm: $75/sf design and permits, loan terms, 5% commission. Figures are from public portals and city records, not MLS.",
]
for k, t in enumerate(notes, ar + 2):
    cc = db.cell(row=k, column=1, value=t); cc.font = sub if k == ar + 2 else small
    db.merge_cells(start_row=k, start_column=1, end_row=k, end_column=len(short) + 1)
    cc.alignment = Alignment(wrap_text=True, vertical="top"); db.row_dimensions[k].height = 14 if k == ar + 2 else 26
db.freeze_panes = "B6"
db.page_setup.orientation = "landscape"; db.page_setup.fitToWidth = 1; db.page_setup.fitToHeight = 1; db.sheet_properties.pageSetUpPr.fitToPage = True
db.sheet_view.showGridLines = False

# ================= Sold Comps =================
NBN = NB
cp = wb.create_sheet("Sold Comps")
cp['A1'] = "Qualifying closed sales (houses built 2010+, 2,000+ sf, sold in the last 24 months, not burned lots). Kept current by update_comps.py."; cp['A1'].font = small
cp.append(["Sold", "Address", "Neighborhood", "Price", "Living sf", "$/sf", "Built", "Fire damage (county)", "Source"])
for c in cp[2]: c.font = hf; c.fill = PatternFill("solid", start_color=NAVY)
for r in COMPS_LOG:
    cp.append([r["sold_date"], r["address"], NBN.get(r["nbhd"], r["nbhd"]), float(r["price"]), float(r["sqft"]), float(r["psf"]), int(float(r["year_built"])), r["damage"], (r["source"] + (" MLS " + r["mls"] if r["mls"] else ""))])
cp.append([])
cp.append(["New builds for sale (asking prices, not sales)"]); cp.cell(row=cp.max_row, column=1).font = sub
cp.append(["Listed", "Address", "Neighborhood", "Asking", "Living sf", "$/sf", "Built", "Status", "Source"])
for c in cp[cp.max_row]: c.font = hf; c.fill = PatternFill("solid", start_color=NAVY)
for r in ASKING:
    cp.append([r["listed"], r["address"], r["neighborhood"], float(r["asking"]), float(r["sqft"]), round(float(r["asking"]) / float(r["sqft"])), r["built"], r["status"], r["source"]])
for row in cp.iter_rows(min_row=3):
    for c in row: c.font = body
    if isinstance(row[3].value, (int, float)):
        row[3].number_format = money0; row[4].number_format = num; row[5].number_format = money0
for k, w in enumerate([11, 26, 18, 13, 9, 8, 7, 22, 34], 1): cp.column_dimensions[L(k)].width = w
cp.freeze_panes = "A3"


wb.active = 0
wb.save(os.path.join(HERE, "Palisades_Lot_Screen.xlsx"))
