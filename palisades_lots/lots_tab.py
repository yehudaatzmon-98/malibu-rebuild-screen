"""Markdown for the data room's "All lots" tab: the whole lot screen, refreshed daily. Usage: python lots_tab.py YYYY-MM-DD"""
import csv, os, sys
from collections import defaultdict
from datetime import date, datetime, timedelta
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__)); D = os.path.join(HERE, "data")
asof = sys.argv[1] if len(sys.argv) > 1 else date.today().isoformat()
ws = openpyxl.load_workbook(os.path.join(HERE, "Palisades_Lot_Screen.xlsx"), data_only=True)["All Lots"]
h = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}
short = [l.strip() for l in open(os.path.join(D, "short.txt")) if l.strip()]
M = lambda x: f"${x / 1e6:.2f}M"
lots = []
for r in range(3, ws.max_row + 1):
    g = lambda k: ws.cell(r, h[k]).value
    if not g("Address"): continue
    lots.append(dict(a=g("Address"), nb=g("Neighborhood"), ask=g("Asking price") or 0, sf=g("New house size (sf)") or 0, how=g("How sized"),
                     psf=g("Sale price per sf") or 0, w=g("Worst: profit") or 0, m=g("Profit") or 0, b=g("Best: profit") or 0,
                     mg=g("Margin (profit ÷ cost)") or 0, mx=g("Most we should pay for the lot") or 0, ap=g("City-approved plans?"), notes=g("Notes") or ""))
ok = sorted([l for l in lots if l["mg"] >= 0.15], key=lambda l: -l["mg"])
rest = sorted([l for l in lots if l["mg"] < 0.15], key=lambda l: (l["ask"] / l["mx"]) if l["mx"] > 0 else 99)
flag = lambda l: (" ★" if l["a"] in short else "") + (" ⚠" if "prior home size unknown" in l["notes"] else "")

out = ["# All lots", "",
       f"Updated {asof} from every active land listing in 90272 on Redfin. {len(lots)} lots are in the screen; "
       f"{len(ok)} clear a 15% margin at asking in the medium case. This tab replaces the spreadsheet: it is rebuilt "
       f"every morning from the same model, so this link is always current.", "",
       "★ = on our shortlist. ⚠ = no prior home on record, so the house size is a zoning guess; check before trusting the numbers.", ""]

# What changed, last 14 days
cut = (datetime.strptime(asof, "%Y-%m-%d").date() - timedelta(days=14)).isoformat()
by = defaultdict(list)
if os.path.exists(os.path.join(D, "lots_changes.csv")):
    for c in csv.DictReader(open(os.path.join(D, "lots_changes.csv"))):
        if c["date"] >= cut:
            if c["change"] == "price":
                by[c["date"]].append(f"{c['address']}: price {M(float(c['old']))} → {M(float(c['new']))}")
            elif c["change"] == "new listing":
                by[c["date"]].append(f"{c['address']}: new listing at {M(float(c['new']))}")
            else:
                by[c["date"]].append(f"{c['address']}: {c['change']}")
for r in csv.DictReader(open(os.path.join(D, "comps_log.csv"))):
    if r.get("added_on", "") >= cut and r["qualifies"] == "Y" and r["sold_date"] >= cut:
        nb = "new build" if int(float(r["year_built"])) >= 2025 else f"built {int(float(r['year_built']))}"
        by[r["added_on"]].append(f"{r['address']}: home sold {r['sold_date']} for {M(float(r['price']))} (${int(float(r['psf'])):,}/sf, {nb}); added to comps")
out += ["## What changed (last 14 days)", ""]
if by:
    for d in sorted(by, reverse=True):
        out += [f"**{d}**"] + [f"- {x}" for x in by[d]] + [""]
else:
    out += ["No changes yet.", ""]

out += ["## Lots that clear 15% at asking", "",
        "| Lot | Neighborhood | Ask | House (sf, basis) | Sale $/sf (medium) | Profit: worst / medium / best | Margin | Max offer | Approved plans |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- |"]
for l in ok:
    out.append(f"| {l['a']}{flag(l)} | {l['nb']} | {M(l['ask'])} | {int(l['sf']):,} ({l['how'].lower()}) | ${int(l['psf']):,} | "
               f"{M(l['w'])} / {M(l['m'])} / {M(l['b'])} | {l['mg']:.0%} | {M(l['mx'])} | {'Yes' if l['ap'] == 'Yes' else ''} |")
out += ["", "## All other lots (closest to working first)", "",
        "| Lot | Neighborhood | Ask | Max offer | Ask ÷ max offer | Medium margin | Approved plans |", "| --- | --- | --- | --- | --- | --- | --- |"]
for l in rest:
    ratio = f"{l['ask'] / l['mx']:.0%}" if l["mx"] > 0 else "never works"
    out.append(f"| {l['a']}{flag(l)} | {l['nb']} | {M(l['ask'])} | {M(max(l['mx'], 0))} | {ratio} | {l['mg']:.0%} | {'Yes' if l['ap'] == 'Yes' else ''} |")
out += ["", "## How the numbers work", "",
        "- Sale price per sf by neighborhood comes from closed sales (see the Comps tracker tab). Worst case = today's comps, medium = +10%, best = +20%.",
        "- House size: city-approved plans where they exist; otherwise the larger-profit of a 110% rebuild of the old home (2-year project) or 40% of lot area up to 8,000 sf (3-year project).",
        "- Costs: construction $700/sf fixed; design, permits and fees $75/sf, or a 5% reserve on the lot price where plans are approved; construction loan 65% of cost at 10% with 1.5 points (60% average draw); property tax 1.2% of land per year; 5% commission, 0.56% transfer tax and Measure ULA (4% over $5.4M, 5.5% over $10.9M).",
        "- Max offer = the most we can pay for the lot and still earn 15% on total cost in the medium case. Profit is before any investor/sponsor split.",
        "- Source: Redfin active land listings, LA County parcel data (old home size), LADBS open data (approved plans)."]
print("\n".join(out))
