"""Snapshot every lot's key numbers from the recalculated workbook, or compare two snapshots.

  python snapshot.py > snap.json            # take a snapshot
  python snapshot.py --diff before.json after.json   # markdown report of what changed
"""
import json, os, sys

HERE = os.path.dirname(os.path.abspath(__file__))


def take():
    import openpyxl
    ws = openpyxl.load_workbook(os.environ.get("WB") or os.path.join(HERE, "Palisades_Lot_Screen.xlsx"), data_only=True)["All Lots"]
    h = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}
    out = {}
    for r in range(3, ws.max_row + 1):
        g = lambda k: ws.cell(r, h[k]).value
        if not g("Address"): continue
        out[g("Address")] = dict(nbhd=g("Neighborhood"), ask=g("Asking price"), sf=g("New house size (sf)"), how=g("How sized"),
                                 margin=round(g("Margin (profit ÷ cost)") or 0, 3), profit=round(g("Profit") or 0),
                                 worst=round(g("Worst: profit") or 0), max_offer=round(g("Most we should pay for the lot") or 0),
                                 approved=g("City-approved plans?"), notes=g("Notes"))
    return out


def diff(a, b, short):
    M = lambda x: f"${x / 1e6:.2f}M"
    lines = []
    clear_b = {k for k, v in b.items() if v["margin"] >= 0.15}
    clear_a = {k for k, v in a.items() if v["margin"] >= 0.15}
    newly = sorted(clear_b - clear_a, key=lambda k: -b[k]["margin"])
    dropped = sorted(k for k in clear_a - clear_b if k in b)
    gone = sorted(set(a) - set(b)); new = sorted(set(b) - set(a))
    cuts = sorted([k for k in b if k in a and b[k]["ask"] != a[k]["ask"]], key=lambda k: b[k]["ask"] / a[k]["ask"])
    lines += [f"Lots in the screen: {len(b)}. Clearing 15% at asking (medium case): {len(clear_b)} (was {len(clear_a)})."]
    if cuts:
        lines += ["", "| Price change | Old ask | New ask | Medium margin | Max offer | Clears 15%? |", "| --- | --- | --- | --- | --- | --- |"]
        lines += [f"| {k}{' (shortlist)' if k in short else ''} | {M(a[k]['ask'])} | {M(b[k]['ask'])} | {b[k]['margin']:.0%} | {M(b[k]['max_offer'])} | {'yes' if b[k]['margin'] >= 0.15 else 'no'} |" for k in cuts]
    if newly:
        lines += ["", "Newly clearing 15%:"] + [f"- {k} ({b[k]['nbhd']}): ask {M(b[k]['ask'])}, margin {b[k]['margin']:.0%}, medium profit {M(b[k]['profit'])}, worst {M(b[k]['worst'])}, max offer {M(b[k]['max_offer'])}; size {b[k]['sf']:,} sf ({b[k]['how']}){' — ' + b[k]['notes'] if b[k]['notes'] else ''}{' **REVIEW: no prior home on record, so the size is a pure zoning guess**' if 'prior home size unknown' in (b[k]['notes'] or '') else ''}" for k in newly]
    if dropped:
        lines += ["", "No longer clearing 15%:"] + [f"- {k}: margin {b[k]['margin']:.0%}" for k in dropped]
    if new:
        lines += ["", "New listings added: " + ", ".join(f"{k} ({M(b[k]['ask'])}, {b[k]['margin']:.0%})" for k in new)]
    if gone:
        lines += ["", "No longer listed: " + ", ".join(gone) + (" — includes a shortlisted lot!" if set(gone) & set(short) else "")]
    if not (cuts or newly or dropped or new or gone):
        lines = ["No changes in the lot list today."]
    return "\n".join(lines)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--diff":
        short = [l.strip() for l in open(os.path.join(HERE, "data", "short.txt")) if l.strip()]
        print(diff(json.load(open(sys.argv[2])), json.load(open(sys.argv[3])), short))
    else:
        print(json.dumps(take(), indent=0))
