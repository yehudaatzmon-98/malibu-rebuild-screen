"""Shared deal math for the investor materials. Reads the current workbook so every number matches the model.

Two financing models (Tal, 10/6 call):
  "loan"  - Construction loan for 65% of total cost (land + construction + design), closed with the purchase.
  "cash"  - Land bought with cash; the construction loan covers construction + design only, capped at 65% of
            total cost. Smaller loan, less interest, more cash in. Tal: buying with a combined purchase +
            construction loan takes about 2 months per lot and is hard to close early on.
"""
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
_wb = openpyxl.load_workbook(os.path.join(HERE, "Palisades_Lot_Screen.xlsx"), data_only=True)
_a = _wb["Assumptions"]
A = {}
r = 4
while _a.cell(r, 1).value:
    A[_a.cell(r, 1).value] = _a.cell(r, 2).value; r += 1
_lt = _wb["All Lots"]; _h = {_lt.cell(2, c).value: c for c in range(1, _lt.max_column + 1)}
LOTS = {}
for r in range(3, _lt.max_row + 1):
    if _lt.cell(r, _h["Address"]).value:
        g = {k: _lt.cell(r, c).value for k, c in _h.items()}
        LOTS[g["Address"]] = g
PLAN = [l.strip() for l in open(os.path.join(HERE, "data", "plan10.txt")) if l.strip()]
COSTS = [A["Construction cost ($/sf, fixed)"], A["Sensitivity: lower build cost A ($/sf)"], A["Sensitivity: lower build cost B ($/sf)"]]
CASES = [("Worst", "Worst: sale $/sf"), ("Medium", "Medium: sale $/sf"), ("Best", "Best: sale $/sf")]
LTC = A["Construction loan (% of cost)"]


def ula(s):
    return A["ULA tier 2 rate"] if s >= A["ULA tier 2 threshold"] else (A["ULA tier 1 rate"] if s >= A["ULA tier 1 threshold"] else 0)


def calc(g, case="Medium", cost=None, fin="loan"):
    """Full build-up for one lot. Returns a dict of every line."""
    cost = COSTS[0] if cost is None else cost
    psf = g[dict(CASES)[case]]
    sf, y, land = g["New house size (sf)"], g["Project length (yrs)"], g["Asking price"]
    approved = g["City-approved plans?"] == "Yes"
    construction = sf * cost
    design = A["Reserve — lots WITH city-approved plans (% of asking price)"] * land if approved else A["Design, permits & fees — lots WITHOUT approved plans ($/sf)"] * sf
    base = land + construction + design
    loan = LTC * base if fin == "loan" else min(LTC * base, construction + design)
    interest = loan * (A["Loan rate"] * y * A["Avg % of loan drawn"] + A["Loan points"])
    tax = land * A["Property tax (% of land / yr)"] * y
    sale = sf * psf
    comm = sale * (A["Sale commission"] + A["City + county transfer tax"])
    u = sale * ula(sale)
    total = base + interest + tax + comm + u
    profit = sale - total
    cash_in = base - loan + interest + tax
    return dict(land=land, sf=sf, years=y, psf=psf, construction=construction, design=design, base=base, loan=loan,
                interest=interest, tax=tax, sale=sale, comm=comm, ula=u, total=total, profit=profit,
                margin=profit / total, cash_in=cash_in, roc=profit / cash_in)


def portfolio(addrs=None, case="Medium", cost=None, fin="loan"):
    addrs = PLAN if addrs is None else addrs
    rows = [calc(LOTS[a], case, cost, fin) for a in addrs if a in LOTS]
    t = {k: sum(r[k] for r in rows) for k in ("land", "construction", "design", "base", "loan", "interest", "tax", "sale", "comm", "ula", "total", "profit", "cash_in")}
    t["margin"] = t["profit"] / t["total"]; t["roc"] = t["profit"] / t["cash_in"]; t["n"] = len(rows)
    t["years"] = sum(r["years"] for r in rows) / len(rows)
    return t


def selfcheck():
    """Financing 'loan' at the base cost must equal the workbook exactly."""
    for a in PLAN:
        g = LOTS[a]
        for case, key in CASES:
            p = calc(g, case)["profit"]
            wb = {"Worst": g["Worst: profit"], "Medium": g["Medium: profit"], "Best": g["Best: profit"]}[case]
            assert abs(p - wb) < 5, (a, case, p, wb)
    return True


if __name__ == "__main__":
    selfcheck()
    for fin in ("loan", "cash"):
        for cost in COSTS:
            row = [portfolio(case=c, cost=cost, fin=fin) for c, _ in CASES]
            print(fin, cost, f"cash in {row[0]['cash_in']/1e6:.1f}M", " | ".join(f"{c}: profit {t['profit']/1e6:.1f}M margin {t['margin']:.0%} on-cash {t['roc']:.0%}" for (c, _), t in zip(CASES, row)))
