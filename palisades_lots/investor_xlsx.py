"""Build the investor workbook (uploaded to Google Drive as a Google Sheet). Live formulas; inputs from deal.py's model.
No deal terms, no max offers. Usage: python investor_xlsx.py YYYY-MM-DD [out.xlsx]"""
import csv, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL
from deal import A, LOTS, PLAN, COSTS, selfcheck

HERE = os.path.dirname(os.path.abspath(__file__)); D = os.path.join(HERE, "data")
asof = sys.argv[1]; out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "investor", "Palisades_Investor_Numbers.xlsx")
selfcheck()
F = lambda **k: Font(**{"name": "Arial", "size": 10, **k})
BLUE = F(color="0000FF"); BOLD = F(bold=True); HF = F(bold=True, color="FFFFFF"); GREEN = F(color="008000")
HFILL = PatternFill("solid", fgColor="13283A"); BAND = PatternFill("solid", fgColor="EDE6DA"); KEY = PatternFill("solid", fgColor="FFF2CC")
USD = '$#,##0;-$#,##0;"–"'; MM = '$#,##0.00,,"M";-$#,##0.00,,"M";"–"'; PCT = '0%'; PCT1 = '0.0%'
NOTES = {r["address"]: r["note"] for r in csv.DictReader(open(os.path.join(D, "investor_notes.csv")))}
nbc = {r["code"]: r for r in csv.DictReader(open(os.path.join(D, "nbhd_base.csv")))}
CODE = {"Riviera": "RIV", "Huntington": "HUNT", "Bluffs / El Medio": "BLUFF"}
SEL = [a for a in PLAN if a in LOTS]
wb = openpyxl.Workbook()


def hdr(ws, r, vals, c0=1):
    for i, v in enumerate(vals):
        c = ws.cell(r, c0 + i, v); c.font = HF; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1): ws.column_dimensions[CL(i)].width = w


# ---- Read me
rm = wb.active; rm.title = "Read me"
lines = [("Rebuilding the Palisades — lots and numbers", F(bold=True, size=14)),
         (f"Updated {asof}. Project-level figures for the ten lots, before any split between investors and sponsors. A discussion draft, not an offer to sell securities.", None),
         ("", None),
         ("How to use it", BOLD),
         ("Summary: the ten lots and the project totals. Lot breakdown: every cost line for each lot. Calc: every lot in every case, build cost and financing (filterable).", None),
         ("Inputs (blue) drive everything: change one and every number recalculates. Lots holds each lot's facts. Market evidence holds the closed sales and new homes for sale.", None),
         ("", None),
         ("The cases", BOLD),
         ("Worst: no premium for a new home over the price each neighborhood uses (Inputs). Medium: +10%. Best: +20%.", None),
         ("Closed-sale medians: each neighborhood priced at the median of its recent qualifying closed sales instead (lower than the price used in the Bluffs and Huntington).", None),
         ("", None),
         ("Financing", BOLD),
         ("A: one construction loan for 65% of land + construction + design, closed with the purchase.", None),
         ("B: land bought with cash; the loan covers construction and design only, up to 65% of total cost.", None),
         ("", None),
         ("Still to confirm: the builder's $700/sf as a fixed price in writing; loan terms; house sizes on lots without plans (architect); closed sales against title; that approved plans transfer with each sale.", None),
         ("Sources: listings from Redfin; closed sales from an MLS export plus Redfin and Compass records; LA County Assessor; LADBS permit records. Listing claims are the sellers' words unless marked as checked with the city.", None)]
for i, (t, f) in enumerate(lines, 1):
    c = rm.cell(i, 1, t); c.font = f or F(); c.alignment = Alignment(wrap_text=True, vertical="top")
rm.column_dimensions["A"].width = 120

# ---- Inputs
ip = wb.create_sheet("Inputs")
ip["A1"] = "Inputs (blue = change me; every number in the workbook updates)"; ip["A1"].font = F(bold=True, size=12)
hdr(ip, 3, ["Input", "Value", "Note"])
rows = [("Construction cost ($/sf)", COSTS[0], USD, "Builder's quote; not yet fixed in writing"),
        ("Lower build cost scenario 1 ($/sf)", COSTS[1], USD, "Tal's estimate"),
        ("Lower build cost scenario 2 ($/sf)", COSTS[2], USD, "Reported locally (secondhand, not verified)"),
        ("Design, permits and fees ($/sf, no approved plans)", A["Design, permits & fees — lots WITHOUT approved plans ($/sf)"], USD, "Placeholder until architect quotes"),
        ("Reserve when plans are approved (% of lot price)", A["Reserve — lots WITH city-approved plans (% of asking price)"], PCT1, "Replaces design costs"),
        ("Construction loan (% of cost)", A["Construction loan (% of cost)"], PCT1, "To confirm with a lender"),
        ("Loan rate", A["Loan rate"], PCT1, "Placeholder"),
        ("Loan points", A["Loan points"], PCT1, "Placeholder"),
        ("Average share of loan drawn", A["Avg % of loan drawn"], PCT1, "Loans are drawn as the house is built"),
        ("Property tax (% of lot price per year)", A["Property tax (% of land / yr)"], PCT1, "About 1.2% of purchase price"),
        ("Sale commission", A["Sale commission"], PCT1, "Tal lists the homes; includes the buyer's agent"),
        ("City + county transfer tax", A["City + county transfer tax"], '0.00%', "LA City 0.45% + County 0.11%"),
        ("ULA tier 1 threshold", A["ULA tier 1 threshold"], USD, "Measure ULA"),
        ("ULA tier 1 rate", A["ULA tier 1 rate"], PCT1, "Measure ULA"),
        ("ULA tier 2 threshold", A["ULA tier 2 threshold"], USD, "Measure ULA"),
        ("ULA tier 2 rate", A["ULA tier 2 rate"], PCT1, "Measure ULA"),
        ("Ocean-view premium ($/sf)", A["Ocean-view premium ($/sf added to the neighborhood price)"], USD, "Where the MLS listing shows an ocean view"),
        ("Medium case premium", A["MEDIUM case: premium over today's comps"], PCT1, "For a brand-new home"),
        ("Best case premium", A["BEST case: premium over today's comps"], PCT1, "As the neighborhood recovers")]
N = {}
for i, (k, v, f, note) in enumerate(rows, 4):
    ip.cell(i, 1, k); c = ip.cell(i, 2, v); c.font = BLUE; c.number_format = f; ip.cell(i, 3, note)
    N[k] = f"Inputs!$B${i}"
r0 = 4 + len(rows) + 1
ip.cell(r0, 1, "Sale price per sf by neighborhood").font = BOLD
hdr(ip, r0 + 1, ["Neighborhood", "Price used ($/sf)", "Closed-sale median ($/sf)", "How the price was set"])
NBR = {}
for j, (nb, code) in enumerate(CODE.items()):
    r = r0 + 2 + j; row = nbc[code]
    med = float(row["rule_psf"]) if row["rule_n"] not in ("", "0") else float(row["applied_psf"])
    ip.cell(r, 1, nb)
    c = ip.cell(r, 2, float(row["applied_psf"])); c.font = BLUE; c.number_format = USD
    c = ip.cell(r, 3, med); c.font = BLUE; c.number_format = USD
    ip.cell(r, 4, f"Median of {row['rule_n']} qualifying sales" if row["status"].startswith("rule") else f"Judgment, above the median of {row['rule_n']} qualifying sales")
NBRNG = (f"Inputs!$A${r0+2}:$A${r0+4}", f"Inputs!$B${r0+2}:$B${r0+4}", f"Inputs!$C${r0+2}:$C${r0+4}")
widths(ip, [48, 18, 22, 48])

# ---- Lots
lt = wb.create_sheet("Lots")
lt["A1"] = f"The ten lots (asking prices and facts as of {asof}; blue = input)"; lt["A1"].font = F(bold=True, size=12)
hdr(lt, 3, ["Lot", "Neighborhood", "Asking price", "Lot size (sf)", "New house (sf)", "How sized", "Years, purchase to sale", "City-approved plans?", "Ocean view (MLS)?", "Facts"])
LR = {}
basis = {"Approved plans": "City-approved plans", "110% old house": "Like-for-like rebuild at 110%", "Zoning estimate": "Estimate from lot area"}
for i, a in enumerate(SEL, 4):
    g = LOTS[a]; LR[a] = i
    vals = [a, g["Neighborhood"], g["Asking price"], g["Lot size (sf)"], g["New house size (sf)"], basis.get(g["How sized"], g["How sized"]),
            g["Project length (yrs)"], g["City-approved plans?"], g.get("Ocean view (MLS)") or "No", NOTES.get(a, "")]
    for j, v in enumerate(vals, 1):
        c = lt.cell(i, j, v)
        if 3 <= j <= 9: c.font = BLUE
        c.alignment = Alignment(wrap_text=(j == 10), vertical="top")
    lt.cell(i, 3).number_format = USD; lt.cell(i, 4).number_format = '#,##0'; lt.cell(i, 5).number_format = '#,##0'
widths(lt, [22, 18, 14, 12, 13, 26, 12, 12, 12, 90])
lt.freeze_panes = "B4"

# ---- Calc
cs = wb.create_sheet("Calc")
cs["A1"] = "Every lot × case × build cost × financing. All formulas."; cs["A1"].font = BOLD
H = ["Lot", "Case", "Build cost ($/sf)", "Financing", "Sale $/sf", "Sale price", "Land", "Construction", "Design or reserve", "Land + build + design",
     "Loan", "Interest and points", "Property tax", "Commission and transfer tax", "ULA tax", "Total cost", "Profit", "Margin", "Cash needed"]
hdr(cs, 3, H)
CASES = [("Worst", 0, "B"), ("Medium", N["Medium case premium"], "B"), ("Best", N["Best case premium"], "B"), ("Closed-sale medians", 0, "C")]
COSTREF = [N["Construction cost ($/sf)"], N["Lower build cost scenario 1 ($/sf)"], N["Lower build cost scenario 2 ($/sf)"]]
r = 4
for a in SEL:
    lr = LR[a]; L_ = lambda c: f"Lots!${c}${lr}"
    for case, prem, col in CASES:
        for cref in COSTREF:
            for fin in ("A", "B"):
                pcol = NBRNG[1] if col == "B" else NBRNG[2]
                f = {1: f"={L_('A')}", 2: case, 3: f"={cref}", 4: fin,
                     5: f'=(INDEX({pcol},MATCH({L_("B")},{NBRNG[0]},0))+IF({L_("I")}="Yes",{N["Ocean-view premium ($/sf)"]},0))*(1+{prem})',
                     6: f"={L_('E')}*E{r}", 7: f"={L_('C')}", 8: f"={L_('E')}*C{r}",
                     9: f'=IF({L_("H")}="Yes",{N["Reserve when plans are approved (% of lot price)"]}*G{r},{N["Design, permits and fees ($/sf, no approved plans)"]}*{L_("E")})',
                     10: f"=G{r}+H{r}+I{r}",
                     11: f'=IF(D{r}="A",{N["Construction loan (% of cost)"]}*J{r},MIN({N["Construction loan (% of cost)"]}*J{r},H{r}+I{r}))',
                     12: f"=K{r}*({N['Loan rate']}*{L_('G')}*{N['Average share of loan drawn']}+{N['Loan points']})",
                     13: f"=G{r}*{N['Property tax (% of lot price per year)']}*{L_('G')}",
                     14: f"=F{r}*({N['Sale commission']}+{N['City + county transfer tax']})",
                     15: f"=F{r}*IF(F{r}>={N['ULA tier 2 threshold']},{N['ULA tier 2 rate']},IF(F{r}>={N['ULA tier 1 threshold']},{N['ULA tier 1 rate']},0))",
                     16: f"=J{r}+L{r}+M{r}+N{r}+O{r}", 17: f"=F{r}-P{r}", 18: f"=IF(P{r}=0,0,Q{r}/P{r})", 19: f"=J{r}-K{r}+L{r}+M{r}"}
                for c, v in f.items():
                    cell = cs.cell(r, c, v); cell.number_format = PCT if c == 18 else (USD if c >= 5 or c == 3 else "General")
                    if c in (1, 3, 7): cell.font = GREEN
                r += 1
LAST = r - 1
widths(cs, [22, 18, 12, 10] + [14] * 15)
cs.freeze_panes = "B4"; cs.auto_filter.ref = f"A3:S{LAST}"
rg = lambda c: f"Calc!${c}$4:${c}${LAST}"
S = lambda col, lot, case, cost, fin: f'SUMIFS({rg(col)},{rg("A")},{lot},{rg("B")},{case},{rg("C")},{cost},{rg("D")},{fin})'

# ---- Summary
sm = wb.create_sheet("Summary", 1)
sm["A1"] = "Rebuilding the Palisades — the ten lots"; sm["A1"].font = F(bold=True, size=14)
sm["A2"] = f"Updated {asof}. Construction at the Inputs build cost; financing A unless noted. Project figures, before any split."; sm["A2"].font = F(italic=True)
hdr(sm, 4, ["Lot", "Neighborhood", "Asking price", "New house (sf)", "How sized", "Sale price (medium)", "Profit: worst", "Profit: medium", "Profit: best", "Profit: closed-sale medians", "Margin (medium)"])
for i, a in enumerate(SEL, 5):
    lr = LR[a]
    for j, c in enumerate("ABCEF", 1):
        cell = sm.cell(i, j, f"=Lots!{c}{lr}"); cell.font = GREEN
    sm.cell(i, 3).number_format = USD; sm.cell(i, 4).number_format = '#,##0'
    sm.cell(i, 6, "=" + S("F", f"$A{i}", '"Medium"', COSTREF[0], '"A"')).number_format = MM
    for j, case in zip(range(7, 11), ("Worst", "Medium", "Best", "Closed-sale medians")):
        sm.cell(i, j, "=" + S("Q", f"$A{i}", f'"{case}"', COSTREF[0], '"A"')).number_format = MM
    sm.cell(i, 11, f"=H{i}/" + S("P", f"$A{i}", '"Medium"', COSTREF[0], '"A"')).number_format = PCT
tr = 5 + len(SEL)
sm.cell(tr, 1, f"All {len(SEL)}").font = BOLD
for j in (3, 6, 7, 8, 9, 10):
    c = sm.cell(tr, j, f"=SUM({CL(j)}5:{CL(j)}{tr-1})"); c.font = BOLD; c.number_format = MM if j != 3 else MM
c = sm.cell(tr, 11, f'=H{tr}/SUMIFS({rg("P")},{rg("B")},"Medium",{rg("C")},{COSTREF[0]},{rg("D")},"A")'); c.font = BOLD; c.number_format = PCT
for j in range(1, 12): sm.cell(tr, j).fill = BAND

r = tr + 3
CASEL = ["Worst", "Medium", "Best", "Closed-sale medians"]
LINES = [("Land (asking prices)", "G"), ("Construction and design", "HI"), ("Loan interest and points", "L"), ("Property tax", "M"),
         ("Commission and transfer tax", "N"), ("ULA tax", "O"), ("Total cost", "P"), ("Home sales", "F"), ("Profit", "Q"),
         ("Margin on total cost", "%"), ("Loan", "K"), ("Cash needed", "S"), ("Profit ÷ cash needed (whole project)", "R")]
for fin, title in (("A", "Financing A · one loan for land and build"), ("B", "Financing B · land in cash, loan to build")):
    sm.cell(r, 1, title).font = F(bold=True, size=12); r += 1
    hdr(sm, r, ["All ten lots"] + CASEL); r += 1
    first = r
    for lab, col in LINES:
        sm.cell(r, 1, lab)
        for j, case in enumerate(CASEL, 2):
            tot = lambda cc: f'SUMIFS({rg(cc)},{rg("B")},"{case}",{rg("C")},{COSTREF[0]},{rg("D")},"{fin}")'
            if col == "HI": v = f"={tot('H')}+{tot('I')}"
            elif col == "%": v = f"=IF({tot('P')}=0,0,{tot('Q')}/{tot('P')})"
            elif col == "R": v = f"=IF({tot('S')}=0,0,{tot('Q')}/{tot('S')})"
            else: v = f"={tot(col)}"
            c = sm.cell(r, j, v); c.number_format = PCT if col in ("%", "R") else MM
        if lab in ("Total cost", "Profit"):
            for j in range(1, 6): sm.cell(r, j).font = BOLD; sm.cell(r, j).fill = BAND
        r += 1
    r += 2
sm.cell(r, 1, "If construction comes in lower (medium case; worst case alongside)").font = F(bold=True, size=12); r += 1
hdr(sm, r, ["Build cost ($/sf)", "A: cash needed", "A: profit, medium", "A: profit, worst", "B: cash needed", "B: profit, medium", "B: profit, worst"]); r += 1
for cref in COSTREF:
    c = sm.cell(r, 1, f"={cref}"); c.font = GREEN; c.number_format = USD
    for j, (col, case, fin) in enumerate([("S", "Medium", "A"), ("Q", "Medium", "A"), ("Q", "Worst", "A"), ("S", "Medium", "B"), ("Q", "Medium", "B"), ("Q", "Worst", "B")], 2):
        sm.cell(r, j, f'=SUMIFS({rg(col)},{rg("B")},"{case}",{rg("C")},$A{r},{rg("D")},"{fin}")').number_format = MM
    r += 1
widths(sm, [36, 18, 14, 13, 26, 16, 14, 14, 14, 18, 12])
sm.freeze_panes = "B5"

# ---- Lot breakdown
lb = wb.create_sheet("Lot breakdown", 2)
lb["A1"] = "Every cost line for each lot (financing A, build cost from Inputs)"; lb["A1"].font = F(bold=True, size=12)
r = 3
BL = [("Sale price per sf", "E", USD), ("Sale price", "F", MM), ("Land", "G", MM), ("Construction", "H", MM), ("Design or reserve", "I", MM),
      ("Loan interest and points", "L", MM), ("Property tax", "M", MM), ("Commission and transfer tax", "N", MM), ("ULA tax", "O", MM),
      ("Total cost", "P", MM), ("Profit", "Q", MM), ("Margin", "R", PCT), ("Cash needed (financing A)", "S", MM)]
for a in SEL:
    c = lb.cell(r, 1, f"=Lots!A{LR[a]}&\" · \"&Lots!B{LR[a]}"); c.font = F(bold=True, size=12); r += 1
    c = lb.cell(r, 1, f"=Lots!J{LR[a]}"); c.alignment = Alignment(wrap_text=True, vertical="top"); lb.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    lb.row_dimensions[r].height = 42; r += 1
    hdr(lb, r, [""] + CASEL); r += 1
    for lab, col, fmt in BL:
        lb.cell(r, 1, lab)
        for j, case in enumerate(CASEL, 2):
            if col == "R":
                lotref = f"Lots!$A${LR[a]}"; cq = '"' + case + '"'
                p_ = S("P", lotref, cq, COSTREF[0], '"A"'); q_ = S("Q", lotref, cq, COSTREF[0], '"A"')
                v = f"=IF({p_}=0,0,{q_}/{p_})"
            else:
                v = "=" + S(col, f"Lots!$A${LR[a]}", f'"{case}"', COSTREF[0], '"A"')
            lb.cell(r, j, v).number_format = fmt
        if lab in ("Total cost", "Profit"):
            for j in range(1, 6): lb.cell(r, j).font = BOLD; lb.cell(r, j).fill = BAND
        r += 1
    lb.cell(r, 1, "Cash needed with financing B (medium)")
    lb.cell(r, 3, "=" + S("S", f"Lots!$A${LR[a]}", '"Medium"', COSTREF[0], '"B"')).number_format = MM
    r += 2
widths(lb, [36, 16, 16, 16, 20])

# ---- Market evidence
me = wb.create_sheet("Market evidence")
me["A1"] = "Market evidence: closed sales and new homes for sale"; me["A1"].font = F(bold=True, size=12)
me["A2"] = "A sale counts if it is a house in 90272 built 2010+, 2,000+ sf, closed in the last 24 months, and not a burned house sold as land."
hdr(me, 4, ["Closed", "Address", "Neighborhood", "Built", "Size (sf)", "Price", "Per sf"])
names = {"RIV": "Riviera", "HUNT": "Huntington", "BLUFF": "Bluffs / El Medio"}
r = 5
for row in sorted([x for x in csv.DictReader(open(os.path.join(D, "comps_log.csv"))) if x["qualifies"] == "Y" and x["nbhd"] in names], key=lambda x: x["sold_date"], reverse=True):
    for j, v in enumerate([row["sold_date"], row["address"], names[row["nbhd"]], int(float(row["year_built"])), int(float(row["sqft"])), float(row["price"]), float(row["psf"])], 1):
        c = me.cell(r, j, v)
    me.cell(r, 5).number_format = '#,##0'; me.cell(r, 6).number_format = USD; me.cell(r, 7).number_format = USD
    r += 1
r += 1
me.cell(r, 1, "Brand-new homes for sale (asking prices, not sales)").font = BOLD; r += 1
hdr(me, r, ["Address", "Area", "Asking", "Size (sf)", "Asking per sf", "Status"]); r += 1
for row in csv.DictReader(open(os.path.join(D, "asking_newbuilds.csv"))):
    a_, s_ = float(row["asking"]), float(row["sqft"])
    for j, v in enumerate([row["address"], row["neighborhood"], a_, s_, f"=C{r}/D{r}", row["status"]], 1): me.cell(r, j, v)
    me.cell(r, 3).number_format = USD; me.cell(r, 4).number_format = '#,##0'; me.cell(r, 5).number_format = USD
    r += 1
r += 1
me.cell(r, 1, "What happened after other fires").font = BOLD; r += 1
hdr(me, r, ["Fire", "What prices did", "Source"]); r += 1
for row in [("Woolsey, Malibu (2018)", "$/sf up 5% inside the burn area vs. 27% outside over 3 years; 29 homes rebuilt and none listed at 30 months", "Redfin; Malibu Times"),
            ("Tubbs, Santa Rosa (2017)", "$/sf up 25% inside vs. 35% outside over 3 years", "Redfin"),
            ("Eaton, Altadena (2025)", "3245 Arrowhead Ct rebuild listed at about $1.9M, closed 6/17/2026 at $1.6M; a local developer estimates new homes sell about 20% below no-fire value", "Compass; Homes.com"),
            ("California fires 2001–2015", "Burned areas gained 2–6% vs. neighbors over 1–4 years, peaking around year 3", "Issler et al., UC Berkeley")]:
    for j, v in enumerate(row, 1): me.cell(r, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
me.cell(r + 1, 1, "Across 90272, the median price per sf is down about 17% from a year earlier (Redfin, Aug 2026).")
widths(me, [24, 34, 20, 10, 12, 14, 60])

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and (c.font is None or c.font.name != "Arial"): c.font = F()
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out); print(out)
